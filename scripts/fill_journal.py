"""Remplit le journal de stage ISIMM (planning + taches quotidiennes)."""

import docx

SRC = r"C:\Users\chana\.claude\uploads\234fa801-38fe-4194-a403-650f1720b6dd\b099db6f-Journal_de_stage_ISIMM.docx"
DST = r"C:\Users\chana\cassandra\Journal_de_stage_ISIMM_rempli.docx"

SUJET = (
    "Cassandra - plateforme d'analyse de performance des API pilotee par l'IA : "
    "detection en couches des degradations d'endpoints (baseline saisonniere + Isolation "
    "Forest), attribution a un deploiement recent et explication automatique en langage "
    "naturel, le tout evalue par injection de fautes."
)

PLANNING = [
    ("08/06 - 26/06",
     "Integration et montee en competences : observabilite applicative (Prometheus, Grafana, "
     "OpenTelemetry), conteneurisation Docker, base de series temporelles TimescaleDB. Etat de "
     "l'art de la detection d'anomalies et des solutions APM du marche."),
    ("29/06 - 10/07",
     "Conception et mise en place de la plateforme : specification de conception, services de "
     "demonstration instrumentes, collecte des metriques, baseline saisonniere par endpoint, "
     "detecteur a seuils, machine a etats anti-flapping et banc d'injection de fautes."),
    ("13/07 - 17/07",
     "Industrialisation et enrichissement : conteneurisation du service de detection, "
     "reproductibilite, alerting Slack, explication des incidents par LLM, registre des "
     "deploiements et correlation degradation / deploiement."),
    ("20/07 - 31/07",
     "Detection ML et analyse de cause racine : Isolation Forest gate par la direction, "
     "campagne d'evaluation face aux seuils statiques, graphe de dependances de services, "
     "tableaux de bord, tests automatises et integration continue."),
    ("03/08 - 07/08",
     "Consolidation : documentation d'architecture, synthese des resultats, redaction du "
     "rapport de stage, demonstration finale et restitution a l'equipe."),
]

TACHES = [
    ("08/06/2026", "Accueil et integration chez TALAN Tunisie. Presentation de l'equipe, du cadre du stage et des outils internes. Mise en place du poste de travail et des acces."),
    ("09/06/2026", "Cadrage du sujet avec l'encadrant : detection automatique des degradations de performance des API. Definition du perimetre, des livrables et du planning."),
    ("10/06/2026", "Etat de l'art de l'observabilite applicative : metriques RED, notions de SLI/SLO, limites du seuillage statique et probleme de la fatigue d'alerte."),
    ("11/06/2026", "Etude comparative des solutions APM du marche (Datadog, Dynatrace, New Relic) et positionnement fonctionnel du projet."),
    ("12/06/2026", "Recherche bibliographique sur la detection d'anomalies non supervisee : Isolation Forest, methodes de series temporelles saisonnieres, criteres d'evaluation."),
    ("15/06/2026", "Prise en main de Docker et Docker Compose : conteneurisation, reseaux, volumes, healthchecks et orchestration d'une stack multi-services."),
    ("16/06/2026", "Prise en main de Prometheus : modele de donnees, langage PromQL, mecanisme de scraping, histogrammes et calcul de quantiles."),
    ("17/06/2026", "Prise en main de Grafana : sources de donnees, construction de panneaux, variables de tableau de bord, annotations et provisioning."),
    ("18/06/2026", "Etude d'OpenTelemetry et du Collector : instrumentation des services, chaine receivers / processors / exporters."),
    ("19/06/2026", "POC jetable : premiere stack locale Prometheus + Grafana avec un service instrumente, validation du flux de metriques de bout en bout."),
    ("22/06/2026", "Prise en main de TimescaleDB : hypertables, politiques de retention et de compression, agregats continus."),
    ("23/06/2026", "Modelisation du schema de donnees : endpoints, features par minute, baselines saisonnieres, anomalies, alertes et deploiements."),
    ("24/06/2026", "POC jetable d'un collecteur Python interrogeant l'API Prometheus et ecrivant les features agregees par endpoint en base."),
    ("25/06/2026", "Prise en main de k6 : ecriture d'un premier scenario de charge et definition d'un profil de trafic realiste (journee compressee)."),
    ("26/06/2026", "POC jetable de services FastAPI de demonstration (catalogue, commandes, paiements) instrumentes pour produire du trafic exploitable."),
    ("29/06/2026", "Conception de la baseline saisonniere : decoupage jour-de-semaine x heure, quantiles p10/p50/p90 par endpoint, gestion du demarrage a froid."),
    ("30/06/2026", "Conception de la detection en couches : couche 0 seuils SLO, couche 1 deviation vs bande saisonniere, couche 2 Isolation Forest, principe du direction gating."),
    ("01/07/2026", "Redaction de la specification de conception : architecture cible, schema de donnees, plan d'evaluation. Revue et validation avec l'encadrant."),
    ("02/07/2026", "Initialisation du depot Git et premiere version de la plateforme : docker-compose, services applicatifs, collecteur OpenTelemetry, collecte des metriques."),
    ("03/07/2026", "Implementation du collecteur de features et du job de calcul des baselines saisonnieres par endpoint."),
    ("06/07/2026", "Developpement du detecteur : boucle de scoring toutes les 60 s, seuils SLO statiques et deviation par rapport a la bande saisonniere."),
    ("07/07/2026", "Implementation de la machine a etats anti-flapping (OK, PENDING, FIRING, RESOLVING) pour stabiliser les alertes et eviter les oscillations."),
    ("08/07/2026", "Developpement du banc d'injection de fautes : scenarios latency_step, latency_creep et error_burst, pilotes par fichiers de configuration."),
    ("09/07/2026", "Ajout des scenarios downstream_slow et pool_shrink ; production des fichiers de verite-terrain servant de labels pour l'evaluation."),
    ("10/07/2026", "Mise en place de la chaine de notification Slack et des premiers tableaux de bord de suivi des alertes."),
    ("13/07/2026", "Correction du pipeline de detection : gestion du pending_since, chemins des fichiers de verite-terrain, provisioning automatique de Grafana."),
    ("14/07/2026", "Travaux de reproductibilite : schema SQL complet et versionne, conteneurisation du service de detection, figeage des dependances Python."),
    ("15/07/2026", "Integration d'un LLM (Gemini) pour generer les explications d'incident : contrat JSON strict (resume, cause suspectee, verifications) et repli deterministe."),
    ("16/07/2026", "Enrichissement des alertes Slack avec la cause suspectee et l'explication generee ; banc de comparaison des modeles et mesure des couts."),
    ("17/07/2026", "Registre des deploiements (API FastAPI dediee) et correlation temporelle entre une degradation detectee et un deploiement recent."),
    ("20/07/2026", "Detection ML : construction des features endpoint-relatives (deviations vs baseline, pente robuste Theil-Sen, ratio p99/p50, deltas court terme) et entrainement de l'Isolation Forest."),
    ("21/07/2026", "Hygiene de la baseline : exclusion des fenetres d'injection du calcul des quantiles (p90 passe de 1391 ms a 317 ms). Cycle de vie des artefacts de modele (versionnage horodate, pointeur latest, garde-fou de derive avant promotion)."),
    ("22/07/2026", "Campagne d'evaluation du detecteur en couches face aux seuils statiques : taux de detection, faux positifs par heure, delai de detection. Tuning du seuil de declenchement a 0.60 (F1 maximal) et mise en evidence de l'inertie du parametre contamination."),
    ("23/07/2026", "Tableaux de bord Grafana (RED, detection, alertes), courbes de sensibilite a la magnitude des fautes, tableau de bord de self-observabilite et passe de securite."),
    ("24/07/2026", "Travaux complementaires : heuristique d'alerte precoce par extrapolation Theil-Sen vers le SLO, et comparaison supervise / non supervise (PR-AUC equivalente, ce qui justifie le choix non supervise)."),
    ("27/07/2026", "Analyse de cause racine : table de dependances entre endpoints modelisant le graphe d'appels (DAG), attribution de confiance calibree et protection anti-cycle au niveau de la base."),
    ("28/07/2026", "Traversee multi-hop bornee du graphe de dependances (CTE recursive, profondeur maximale parametrable) et durcissement de l'analyse en mode best-effort."),
    ("29/07/2026", "Suite de tests automatises (105 tests pytest sur la logique pure) et mise en place de l'integration continue. Documentation d'architecture avec diagrammes."),
    ("30/07/2026", "Performance et exploitation : compression et retention TimescaleDB (ratio ~6x), monitoring de la derive du modele en production, exposition de metriques Prometheus natives du detecteur."),
    ("31/07/2026", "Corrections d'exploitation : healthcheck de la base dans Docker Compose, ajustement des seuils SLO du service paiements, corrections des annotations et panneaux des tableaux de bord."),
    ("03/08/2026", "Script de demonstration de bout en bout repetable (trafic, deploiement, injection, detection, attribution) et redaction du guide de demonstration."),
    ("04/08/2026", "Self-observabilite : separation des indicateurs endpoints_with_features et endpoints_scored pour distinguer collecte et scoring."),
    ("05/08/2026", "Consolidation des resultats de la campagne d'evaluation et redaction de la synthese des messages cles a defendre."),
    ("06/08/2026", "Redaction et mise en forme du rapport de stage : generation automatisee du PDF, diagrammes et relecture."),
    ("07/08/2026", "Repetition et execution de la demonstration finale, restitution des travaux a l'equipe et bilan de stage avec l'encadrant."),
]


def set_cell(cell, texte, gras=False):
    """Ecrit dans une cellule en conservant le style du paragraphe existant."""
    para = cell.paragraphs[0]
    for run in list(para.runs):
        run._element.getparent().remove(run._element)
    run = para.add_run(texte)
    run.bold = gras


def main():
    doc = docx.Document(SRC)

    # Sujet de stage (paragraphe libre)
    for para in doc.paragraphs:
        if para.text.strip().startswith("Sujet de stage"):
            para.add_run(" " + SUJET)
            break

    # B - Planning (table 4) : ligne 0 = en-tete
    table_b = doc.tables[4]
    for i, (date, travail) in enumerate(PLANNING, start=1):
        set_cell(table_b.rows[i].cells[0], date, gras=True)
        set_cell(table_b.rows[i].cells[1], travail)

    # C - Taches quotidiennes (table 6) : ligne 0 = en-tete
    table_c = doc.tables[6]
    for i, (date, travail) in enumerate(TACHES, start=1):
        set_cell(table_c.rows[i].cells[0], date, gras=True)
        set_cell(table_c.rows[i].cells[1], travail)

    doc.save(DST)
    print("Ecrit :", DST)
    print("Planning :", len(PLANNING), "lignes / Taches :", len(TACHES), "lignes")


if __name__ == "__main__":
    main()
